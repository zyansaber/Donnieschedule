# -*- coding: utf-8 -*-
"""
Regent RV - Firebase stock_transfer -> SAP validation report

流程：
1. 读取 Firebase Realtime Database /stock_transfer。
2. 提取 chassis、currentLocation、targetLocation、transferType、savedAt。
3. 根据 chassis 反查 MANDT=800、VKORG=3110 的 Sales Order item 0010。
4. 查询 Sales Order 最后有效 PGI：
   - 优先按 601 未被 602 反冲的方式判断；
   - 如果系统字段不支持，回退到 601 - 602 净数量逻辑。
5. 对已 PGI 的 Sales Order 查询最后有效发票，并获取 Invoice-to BP 和名称。
6. 查询最后 Invoice-to BP 的客户主数据最后变更人员（CDHDR，OBJECTCLAS='DEBI'）。
7. 查询公司仓库库存位置：
   - 3211: Perth 0001/0002, Traralgon 0003/0004,
           Launceston 0005/0006, Geelong 0007/0008
   - 3411: Frankston 0002
8. 输出 Excel。

依赖：
    pip install firebase-admin pandas pyodbc openpyxl
"""

import logging
import os
import argparse
import html
import json
import re
import socket
import time
import uuid
from datetime import datetime, timedelta, timezone
from typing import Dict, Iterable, List, Tuple
from urllib.parse import urlencode

import firebase_admin
import pandas as pd
import pyodbc
from firebase_admin import credentials, db
from google.auth.exceptions import RefreshError
from openpyxl.styles import Alignment, Font, PatternFill


# ============================================================
# Firebase
# ============================================================
FIREBASE_CRED_PATH = "firebase-adminsdk.json"
FIREBASE_DB_URL = (
    "https://scheduling-dd672-default-rtdb.asia-southeast1.firebasedatabase.app"
)
FIREBASE_NODE = "stock_transfer"
FIREBASE_ABNORMAL_NODE = "stock_transfer_abnormal"
FIREBASE_AUDIT_NODE = "stock_transfer_audit"
FIREBASE_SYNC_LOCK_NODE = "sap_sync_locks/stock_transfer_worker"
FIREBASE_CONFIG_NODE = "stockTransferWorkflowConfig"
FIREBASE_EMAIL_JOBS_NODE = "email_jobs"


# ============================================================
# SAP HANA
# ============================================================
DSN = (
    "DRIVER={HDBODBC};"
    "SERVERNODE=10.11.2.25:30241;"
    "UID=BAOJIANFENG;"
    "PWD=Xja@2025ABC;"
)

MANDT = "800"
SALES_ORG = "3110"
SALES_ORDER_ITEM = "000010"


# ============================================================
# Worker mode
# ============================================================
WORKER_INTERVAL_SECONDS = 60
WORKER_BATCH_LIMIT = 500
WORKER_LOCK_TTL_SECONDS = 10 * 60
PROCESSING_STALE_SECONDS = 10 * 60
ERROR_RETRY_SECONDS = 10 * 60
ACTIVE_EMAIL_JOB_STATUSES = {"pending", "sending", "retrying", "sent"}
DEFAULT_APP_BASE_URL = "https://schedule-final-tyn6.onrender.com"


# ============================================================
# Warehouse mapping
# ============================================================
WAREHOUSE_MAP: Dict[Tuple[str, str], str] = {
    ("3211", "0001"): "Perth",
    ("3211", "0002"): "Perth",
    ("3211", "0003"): "Traralgon",
    ("3211", "0004"): "Traralgon",
    ("3211", "0005"): "Launceston",
    ("3211", "0006"): "Launceston",
    ("3211", "0007"): "Geelong",
    ("3211", "0008"): "Geelong",
    ("3411", "0002"): "Frankston",
}

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_FILE = os.path.join(
    SCRIPT_DIR,
    "stock_transfer_sap_pgi_invoice_stock_report.xlsx",
)
LOG_FILE = os.path.join(
    SCRIPT_DIR,
    "stock_transfer_sap_pgi_invoice_stock_report.log",
)
RUN_LOG_DIR = os.path.join(SCRIPT_DIR, "stock_transfer_run_logs")


class FirebaseCredentialError(RuntimeError):
    """Raised when the Firebase service-account file cannot authenticate."""


# ============================================================
# Logging
# ============================================================
logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s - %(levelname)s - %(message)s",
    handlers=[
        logging.FileHandler(LOG_FILE, encoding="utf-8"),
        logging.StreamHandler(),
    ],
)
logger = logging.getLogger("StockTransferSAPReport")


def add_once_run_log_handler() -> str:
    os.makedirs(RUN_LOG_DIR, exist_ok=True)
    run_started_at = datetime.now().strftime("%Y%m%d_%H%M%S")
    run_log_file = os.path.join(
        RUN_LOG_DIR,
        f"stock_transfer_once_{run_started_at}.txt",
    )
    handler = logging.FileHandler(run_log_file, encoding="utf-8")
    handler.setLevel(logging.INFO)
    handler.setFormatter(
        logging.Formatter("%(asctime)s - %(levelname)s - %(message)s")
    )
    logging.getLogger().addHandler(handler)
    return run_log_file


# ============================================================
# Helpers
# ============================================================
def normalize_chassis(value) -> str:
    if value is None:
        return ""
    return (
        str(value)
        .strip()
        .upper()
        .replace(" ", "")
        .replace("-", "")
        .replace("_", "")
    )


def clean_text(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def sql_quote_list(values: Iterable[str]) -> str:
    clean_values = []
    for value in values:
        text = clean_text(value).replace("'", "''")
        if text:
            clean_values.append(f"'{text}'")
    return ",".join(clean_values) if clean_values else "''"


def chunk_list(values: List[str], size: int = 300):
    for index in range(0, len(values), size):
        yield values[index:index + size]


SAP_WRITE_KEYWORD_PATTERN = re.compile(
    r"\b(INSERT|UPDATE|DELETE|MERGE|UPSERT|CALL|EXEC|TRUNCATE|DROP|ALTER|CREATE|REPLACE|GRANT|REVOKE)\b",
    re.IGNORECASE,
)


def assert_sap_read_only_sql(sql: str) -> None:
    normalized = sql.strip()
    if not normalized.upper().startswith(("SELECT", "WITH")):
        raise RuntimeError("Blocked non-read-only SAP SQL. Only SELECT/WITH queries are allowed.")
    blocked_keyword = SAP_WRITE_KEYWORD_PATTERN.search(normalized)
    if blocked_keyword:
        raise RuntimeError(f"Blocked SAP SQL containing write keyword: {blocked_keyword.group(1).upper()}")


def hana_query(sql: str) -> pd.DataFrame:
    assert_sap_read_only_sql(sql)
    with pyodbc.connect(DSN) as connection:
        return pd.read_sql(sql, connection)


def ensure_columns(df: pd.DataFrame, columns: List[str]) -> pd.DataFrame:
    result = df.copy()
    for column in columns:
        if column not in result.columns:
            result[column] = ""
    return result


def normalize_sap_text_columns(df: pd.DataFrame, columns: List[str]) -> pd.DataFrame:
    result = df.copy()
    for column in columns:
        if column in result.columns:
            result[column] = result[column].fillna("").astype(str).str.strip()
    return result


def format_sap_date_series(series: pd.Series) -> pd.Series:
    converted = pd.to_datetime(series, errors="coerce")
    return converted.dt.strftime("%Y-%m-%d").fillna("")


def utc_now() -> datetime:
    return datetime.now(timezone.utc)


def utc_now_iso() -> str:
    return utc_now().replace(microsecond=0).isoformat().replace("+00:00", "Z")


def parse_iso_datetime(value) -> datetime:
    text = clean_text(value)
    if not text:
        return datetime.min.replace(tzinfo=timezone.utc)
    try:
        if text.endswith("Z"):
            text = text[:-1] + "+00:00"
        parsed = datetime.fromisoformat(text)
        if parsed.tzinfo is None:
            parsed = parsed.replace(tzinfo=timezone.utc)
        return parsed.astimezone(timezone.utc)
    except ValueError:
        return datetime.min.replace(tzinfo=timezone.utc)


def get_worker_id() -> str:
    return f"{socket.gethostname()}:{os.getpid()}"


def resolve_firebase_cred_path() -> str:
    cred_path = os.environ.get("FIREBASE_CRED_PATH", FIREBASE_CRED_PATH)
    if not os.path.isabs(cred_path):
        script_relative_path = os.path.join(SCRIPT_DIR, cred_path)
        cred_path = (
            script_relative_path
            if os.path.exists(script_relative_path)
            else os.path.abspath(cred_path)
        )
    return cred_path


def get_firebase_credential_summary() -> Dict[str, str]:
    cred_path = resolve_firebase_cred_path()
    try:
        with open(cred_path, "r", encoding="utf-8") as cred_file:
            data = json.load(cred_file)
    except (OSError, json.JSONDecodeError):
        return {"path": cred_path}

    return {
        "path": cred_path,
        "project_id": clean_text(data.get("project_id")),
        "client_email": clean_text(data.get("client_email")),
        "private_key_id": clean_text(data.get("private_key_id")),
    }


def format_firebase_auth_error(error: Exception) -> str:
    summary = get_firebase_credential_summary()
    lines = [
        "Firebase authentication failed before the SAP sync could start.",
        f"Credential file: {summary.get('path', resolve_firebase_cred_path())}",
    ]

    if summary.get("project_id"):
        lines.append(f"Project: {summary['project_id']}")
    if summary.get("client_email"):
        lines.append(f"Service account: {summary['client_email']}")
    if summary.get("private_key_id"):
        lines.append(f"Private key ID: {summary['private_key_id']}")

    message = str(error)
    if "Invalid JWT Signature" in message:
        lines.extend(
            [
                "",
                "Google rejected the service-account JWT signature. The local",
                "firebase-adminsdk.json key is probably stale, revoked, or edited.",
                "Download a fresh service-account JSON key for this Firebase",
                "project and replace the credential file above.",
            ]
        )
    elif "invalid_grant" in message:
        lines.extend(
            [
                "",
                "Google rejected the service-account token grant. Replace the",
                "credential JSON, and also check that this computer's date/time",
                "is correct.",
            ]
        )
    else:
        lines.extend(["", f"Firebase auth error: {message}"])

    return "\n".join(lines)


def is_transfer_syncable(record: dict) -> bool:
    return not (
        record.get("deletedAt")
        or record.get("cancelledAt")
    )


def is_retry_due(record: dict, now_dt: datetime) -> bool:
    retry_after = parse_iso_datetime(record.get("sapSyncRetryAfter"))
    return retry_after <= now_dt


def is_processing_stale(record: dict, now_dt: datetime) -> bool:
    started_at = parse_iso_datetime(record.get("sapSyncStartedAt"))
    return started_at + timedelta(seconds=PROCESSING_STALE_SECONDS) <= now_dt


# ============================================================
# Firebase stock_transfer
# ============================================================
def initialize_firebase() -> None:
    if not firebase_admin._apps:
        cred_path = resolve_firebase_cred_path()
        try:
            cred = credentials.Certificate(cred_path)
        except (OSError, ValueError) as error:
            raise FirebaseCredentialError(format_firebase_auth_error(error)) from error
        firebase_admin.initialize_app(
            cred,
            {"databaseURL": FIREBASE_DB_URL},
        )


def stock_transfer_records_to_df(raw_data: dict) -> pd.DataFrame:
    rows = []
    iterable = raw_data.items() if isinstance(raw_data, dict) else []

    for firebase_key, record in iterable:
        if not isinstance(record, dict):
            continue

        chassis = normalize_chassis(record.get("chassis"))
        if not chassis:
            continue

        rows.append(
            {
                "Firebase Key": str(firebase_key),
                "Chassis": chassis,
                "Transfer Current Location": clean_text(
                    record.get("currentLocation")
                ),
                "Transfer Target Location": clean_text(
                    record.get("targetLocation")
                ),
                "Transfer Type": clean_text(record.get("transferType")),
                "Transfer Saved At": clean_text(record.get("savedAt")),
            }
        )

    columns = [
        "Firebase Key",
        "Chassis",
        "Transfer Current Location",
        "Transfer Target Location",
        "Transfer Type",
        "Transfer Saved At",
    ]
    if not rows:
        return pd.DataFrame(columns=columns)

    result = pd.DataFrame(rows)
    result["Duplicate Chassis in Firebase"] = result.duplicated(
        subset=["Chassis"], keep=False
    ).map({True: "Yes", False: "No"})
    return result


def fetch_stock_transfer() -> pd.DataFrame:
    initialize_firebase()
    raw_data = db.reference(f"/{FIREBASE_NODE}").get() or {}
    return stock_transfer_records_to_df(raw_data)


def should_sync_stock_transfer(record: dict, now_dt: datetime) -> bool:
    if not isinstance(record, dict) or not is_transfer_syncable(record):
        return False

    chassis = normalize_chassis(record.get("chassis"))
    if not chassis:
        return False

    status = clean_text(record.get("sapSyncStatus")).lower()
    if status == "processing" and not is_processing_stale(record, now_dt):
        return False
    if status == "error" and not is_retry_due(record, now_dt):
        return False

    # SAP is the source of truth. Re-sync every non-deleted/non-cancelled
    # stock transfer, including older records that were previously marked done
    # or completed in the workflow, because SAP fields can still change later.
    return True


def get_sync_sort_key(item: tuple, now_dt: datetime) -> tuple:
    _firebase_key, record = item
    status = clean_text(record.get("sapSyncStatus")).lower()

    if status == "pending":
        priority = 0
        date_value = clean_text(record.get("sapSyncRequestedAt"))
    elif status == "error":
        priority = 1
        date_value = clean_text(record.get("sapSyncRetryAfter"))
    elif status == "processing":
        priority = 2
        date_value = clean_text(record.get("sapSyncStartedAt"))
    elif not clean_text(record.get("Sales Order Display")):
        priority = 3
        date_value = clean_text(record.get("createdAt")) or clean_text(record.get("savedAt"))
    else:
        priority = 4
        date_value = clean_text(record.get("sapSyncedAt"))

    parsed_date = parse_iso_datetime(date_value)
    if parsed_date == datetime.min.replace(tzinfo=timezone.utc):
        parsed_date = now_dt - timedelta(days=3650)

    return (
        priority,
        parsed_date,
        clean_text(record.get("createdAt")) or clean_text(record.get("savedAt")),
    )


def fetch_stock_transfer_records_for_sync(limit: int) -> dict:
    initialize_firebase()
    raw_data = db.reference(f"/{FIREBASE_NODE}").get() or {}
    now_dt = utc_now()
    sync_items = [
        (firebase_key, record)
        for firebase_key, record in raw_data.items()
        if should_sync_stock_transfer(record, now_dt)
    ]
    sync_items.sort(key=lambda item: get_sync_sort_key(item, now_dt))
    if limit and limit > 0:
        sync_items = sync_items[:limit]
    return dict(sync_items)


def acquire_worker_lock(worker_id: str) -> bool:
    initialize_firebase()
    lock_ref = db.reference(f"/{FIREBASE_SYNC_LOCK_NODE}")
    lock_token = str(uuid.uuid4())
    now = utc_now()
    expires_at = now + timedelta(seconds=WORKER_LOCK_TTL_SECONDS)

    def update_lock(current):
        current = current if isinstance(current, dict) else {}
        current_expires_at = parse_iso_datetime(current.get("expiresAt"))
        if current.get("lockedBy") and current_expires_at > now:
            return current
        return {
            "lockedBy": worker_id,
            "lockToken": lock_token,
            "lockedAt": utc_now_iso(),
            "expiresAt": expires_at.replace(microsecond=0)
            .isoformat()
            .replace("+00:00", "Z"),
        }

    locked = lock_ref.transaction(update_lock)
    return (
        isinstance(locked, dict)
        and locked.get("lockedBy") == worker_id
        and locked.get("lockToken") == lock_token
    )


def release_worker_lock(worker_id: str) -> None:
    lock_ref = db.reference(f"/{FIREBASE_SYNC_LOCK_NODE}")
    current = lock_ref.get() or {}
    if isinstance(current, dict) and current.get("lockedBy") == worker_id:
        lock_ref.delete()


def claim_stock_transfer_record(firebase_key: str, worker_id: str) -> bool:
    record_ref = db.reference(f"/{FIREBASE_NODE}/{firebase_key}")
    claim_id = str(uuid.uuid4())
    now_dt = utc_now()

    def update_record(current):
        current = current if isinstance(current, dict) else {}
        if not should_sync_stock_transfer(current, now_dt):
            return current

        return {
            **current,
            "sapSyncStatus": "processing",
            "sapSyncStartedAt": utc_now_iso(),
            "sapSyncWorker": worker_id,
            "sapSyncClaimId": claim_id,
            "sapSyncAttempt": int(current.get("sapSyncAttempt") or 0) + 1,
            "sapSyncError": "",
        }

    claimed = record_ref.transaction(update_record)
    return (
        isinstance(claimed, dict)
        and claimed.get("sapSyncStatus") == "processing"
        and claimed.get("sapSyncWorker") == worker_id
        and claimed.get("sapSyncClaimId") == claim_id
    )


# ============================================================
# Chassis -> 3110 Sales Order
# ============================================================
def fetch_sales_orders_by_chassis(chassis_list: List[str]) -> pd.DataFrame:
    columns = [
        "Chassis",
        "Sales Order SAP",
        "Sales Order Display",
        "Sales Order Item",
        "SO Material",
        "SO Material Name",
    ]
    frames = []

    for batch in chunk_list(chassis_list):
        chassis_sql = sql_quote_list(batch)
        sql = f"""
        SELECT DISTINCT
            objk."SERNR" AS "Chassis",
            vbak."VBELN" AS "Sales Order SAP",
            ser02."POSNR" AS "Sales Order Item",
            vbap."MATNR" AS "SO Material",
            COALESCE(NULLIF(vbap."ARKTX", ''), makt."MAKTX", '')
                AS "SO Material Name"
        FROM "SAPHANADB"."SER02" ser02
        INNER JOIN "SAPHANADB"."OBJK" objk
            ON objk."MANDT" = ser02."MANDT"
           AND objk."OBKNR" = ser02."OBKNR"
        INNER JOIN "SAPHANADB"."VBAK" vbak
            ON vbak."MANDT" = ser02."MANDT"
           AND vbak."VBELN" = ser02."SDAUFNR"
        INNER JOIN "SAPHANADB"."VBAP" vbap
            ON vbap."MANDT" = ser02."MANDT"
           AND vbap."VBELN" = ser02."SDAUFNR"
           AND vbap."POSNR" = ser02."POSNR"
        LEFT JOIN "SAPHANADB"."MAKT" makt
            ON makt."MANDT" = vbap."MANDT"
           AND makt."MATNR" = vbap."MATNR"
           AND makt."SPRAS" = 'E'
        WHERE ser02."MANDT" = '{MANDT}'
          AND objk."SERNR" IN ({chassis_sql})
          AND vbak."VKORG" = '{SALES_ORG}'
          AND ser02."POSNR" = '{SALES_ORDER_ITEM}'
        """
        frame = hana_query(sql)
        if not frame.empty:
            frames.append(frame)

    if not frames:
        return pd.DataFrame(columns=columns)

    result = pd.concat(frames, ignore_index=True).drop_duplicates()
    result = normalize_sap_text_columns(
        result,
        [
            "Chassis",
            "Sales Order SAP",
            "Sales Order Item",
            "SO Material",
            "SO Material Name",
        ],
    )
    result["Sales Order Display"] = result["Sales Order SAP"].apply(
        lambda value: value.lstrip("0") or value
    )
    return result[columns]


# ============================================================
# Last valid PGI
# ============================================================
def fetch_last_valid_pgi_unreversed(sales_orders: List[str]) -> pd.DataFrame:
    """
    优先逻辑：找 601，并排除被 602 通过 SMBLN/SJAHR/SMBLP 指向的原始行。
    """
    columns = [
        "Sales Order SAP",
        "SO Is PGI",
        "SO PGI Post Date",
        "SO PGI Month",
        "SO PGI Material Document",
    ]
    frames = []

    for batch in chunk_list(sales_orders):
        so_sql = sql_quote_list(batch)
        sql = f"""
        WITH valid_601 AS (
            SELECT
                pgi."KDAUF" AS "Sales Order SAP",
                pgi."BUDAT_MKPF" AS "PGI Post Date",
                pgi."MBLNR" AS "Material Document",
                pgi."MJAHR" AS "Material Document Year",
                pgi."ZEILE" AS "Material Document Item",
                ROW_NUMBER() OVER (
                    PARTITION BY pgi."KDAUF"
                    ORDER BY
                        pgi."BUDAT_MKPF" DESC,
                        pgi."MJAHR" DESC,
                        pgi."MBLNR" DESC,
                        pgi."ZEILE" DESC
                ) AS rn
            FROM "SAPHANADB"."NSDM_V_MSEG" pgi
            LEFT JOIN "SAPHANADB"."NSDM_V_MSEG" reversal
                ON reversal."MANDT" = pgi."MANDT"
               AND reversal."BWART" = '602'
               AND reversal."SMBLN" = pgi."MBLNR"
               AND reversal."SJAHR" = pgi."MJAHR"
               AND reversal."SMBLP" = pgi."ZEILE"
            WHERE pgi."MANDT" = '{MANDT}'
              AND pgi."BWART" = '601'
              AND pgi."KDAUF" IN ({so_sql})
              AND reversal."MBLNR" IS NULL
        )
        SELECT
            "Sales Order SAP",
            'Yes' AS "SO Is PGI",
            "PGI Post Date" AS "SO PGI Post Date",
            "Material Document" AS "SO PGI Material Document"
        FROM valid_601
        WHERE rn = 1
        """
        frame = hana_query(sql)
        if not frame.empty:
            frames.append(frame)

    if not frames:
        return pd.DataFrame(columns=columns)

    result = pd.concat(frames, ignore_index=True).drop_duplicates(
        subset=["Sales Order SAP"], keep="first"
    )
    result["SO PGI Post Date"] = format_sap_date_series(
        result["SO PGI Post Date"]
    )
    result["SO PGI Month"] = result["SO PGI Post Date"].str.slice(0, 7)
    return ensure_columns(result, columns)[columns]


def fetch_last_valid_pgi_net_fallback(sales_orders: List[str]) -> pd.DataFrame:
    """
    回退逻辑：601数量 - 602数量 > 0，则认为仍有有效PGI；日期取最后601日期。
    """
    columns = [
        "Sales Order SAP",
        "SO Is PGI",
        "SO PGI Post Date",
        "SO PGI Month",
        "SO PGI Material Document",
    ]
    frames = []

    for batch in chunk_list(sales_orders):
        so_sql = sql_quote_list(batch)
        sql = f"""
        WITH movement_summary AS (
            SELECT
                mseg."KDAUF" AS "Sales Order SAP",
                SUM(
                    CASE
                        WHEN mseg."BWART" = '601' THEN COALESCE(mseg."MENGE", 0)
                        WHEN mseg."BWART" = '602' THEN -1 * COALESCE(mseg."MENGE", 0)
                        ELSE 0
                    END
                ) AS "Net PGI Qty",
                MAX(
                    CASE WHEN mseg."BWART" = '601'
                         THEN mseg."BUDAT_MKPF" END
                ) AS "SO PGI Post Date"
            FROM "SAPHANADB"."NSDM_V_MSEG" mseg
            WHERE mseg."MANDT" = '{MANDT}'
              AND mseg."KDAUF" IN ({so_sql})
              AND mseg."BWART" IN ('601', '602')
            GROUP BY mseg."KDAUF"
        )
        SELECT
            "Sales Order SAP",
            CASE WHEN "Net PGI Qty" > 0 THEN 'Yes' ELSE 'No' END AS "SO Is PGI",
            CASE WHEN "Net PGI Qty" > 0 THEN "SO PGI Post Date" ELSE NULL END
                AS "SO PGI Post Date",
            '' AS "SO PGI Material Document"
        FROM movement_summary
        """
        frame = hana_query(sql)
        if not frame.empty:
            frames.append(frame)

    if not frames:
        return pd.DataFrame(columns=columns)

    result = pd.concat(frames, ignore_index=True).drop_duplicates(
        subset=["Sales Order SAP"], keep="first"
    )
    result["SO PGI Post Date"] = format_sap_date_series(
        result["SO PGI Post Date"]
    )
    result["SO PGI Month"] = result["SO PGI Post Date"].str.slice(0, 7)
    return ensure_columns(result, columns)[columns]


def fetch_last_valid_pgi(sales_orders: List[str]) -> pd.DataFrame:
    try:
        result = fetch_last_valid_pgi_unreversed(sales_orders)
        logger.info("PGI query used unreversed 601/602 reference logic.")
        return result
    except Exception as error:
        logger.warning(
            "Unreversed PGI query failed; fallback to net 601-602 logic: %s",
            error,
        )
        return fetch_last_valid_pgi_net_fallback(sales_orders)


# ============================================================
# Last valid invoice + Invoice-to BP
# ============================================================
def fetch_last_valid_invoice(sales_orders: List[str]) -> pd.DataFrame:
    columns = [
        "Sales Order SAP",
        "Last Invoice Number",
        "Last Invoice Date",
        "Invoice-to BP",
        "Invoice-to Name",
    ]
    frames = []

    for batch in chunk_list(sales_orders):
        so_sql = sql_quote_list(batch)
        sql = f"""
        WITH invoice_source AS (
            SELECT DISTINCT
                vbrp."AUBEL" AS "Sales Order SAP",
                vbrk."VBELN" AS "Invoice Number",
                vbrk."FKDAT" AS "Invoice Date",
                COALESCE(
                    NULLIF(vbpa_re."KUNNR", ''),
                    NULLIF(vbpa_rg."KUNNR", ''),
                    NULLIF(vbrk."KUNRG", '')
                ) AS "Invoice-to BP",
                ROW_NUMBER() OVER (
                    PARTITION BY vbrp."AUBEL"
                    ORDER BY
                        vbrk."FKDAT" DESC,
                        vbrk."ERDAT" DESC,
                        vbrk."ERZET" DESC,
                        vbrk."VBELN" DESC
                ) AS rn
            FROM "SAPHANADB"."VBRP" vbrp
            INNER JOIN "SAPHANADB"."VBRK" vbrk
                ON vbrk."MANDT" = vbrp."MANDT"
               AND vbrk."VBELN" = vbrp."VBELN"
            LEFT JOIN "SAPHANADB"."VBPA" vbpa_re
                ON vbpa_re."MANDT" = vbrk."MANDT"
               AND vbpa_re."VBELN" = vbrk."VBELN"
               AND vbpa_re."POSNR" = '000000'
               AND vbpa_re."PARVW" = 'RE'
            LEFT JOIN "SAPHANADB"."VBPA" vbpa_rg
                ON vbpa_rg."MANDT" = vbrk."MANDT"
               AND vbpa_rg."VBELN" = vbrk."VBELN"
               AND vbpa_rg."POSNR" = '000000'
               AND vbpa_rg."PARVW" = 'RG'
            WHERE vbrp."MANDT" = '{MANDT}'
              AND vbrp."AUBEL" IN ({so_sql})
              AND COALESCE(vbrp."AUPOS", '{SALES_ORDER_ITEM}') = '{SALES_ORDER_ITEM}'
              AND COALESCE(vbrk."FKSTO", '') <> 'X'
        )
        SELECT
            source."Sales Order SAP",
            source."Invoice Number" AS "Last Invoice Number",
            source."Invoice Date" AS "Last Invoice Date",
            source."Invoice-to BP",
            COALESCE(kna1."NAME1", '') AS "Invoice-to Name"
        FROM invoice_source source
        LEFT JOIN "SAPHANADB"."KNA1" kna1
            ON kna1."MANDT" = '{MANDT}'
           AND kna1."KUNNR" = source."Invoice-to BP"
        WHERE source.rn = 1
        """
        frame = hana_query(sql)
        if not frame.empty:
            frames.append(frame)

    if not frames:
        return pd.DataFrame(columns=columns)

    result = pd.concat(frames, ignore_index=True).drop_duplicates(
        subset=["Sales Order SAP"], keep="first"
    )
    result = normalize_sap_text_columns(
        result,
        [
            "Sales Order SAP",
            "Last Invoice Number",
            "Invoice-to BP",
            "Invoice-to Name",
        ],
    )
    result["Last Invoice Date"] = format_sap_date_series(
        result["Last Invoice Date"]
    )
    return ensure_columns(result, columns)[columns]


# ============================================================
# Last customer master change user for Invoice-to BP
# ============================================================
def fetch_invoice_bp_last_change(invoice_bps: List[str]) -> pd.DataFrame:
    columns = [
        "Invoice-to BP",
        "Invoice BP Last Changed By",
        "Invoice BP Last Change Date",
        "Invoice BP Last Change Time",
        "Invoice BP Last Change TCode",
    ]
    if not invoice_bps:
        return pd.DataFrame(columns=columns)

    frames = []
    for batch in chunk_list(invoice_bps):
        padded = [clean_text(value).zfill(10) for value in batch]
        bp_sql = sql_quote_list(padded)
        sql = f"""
        WITH ranked_change AS (
            SELECT
                cdhdr."OBJECTID" AS "Invoice-to BP",
                cdhdr."USERNAME" AS "Invoice BP Last Changed By",
                cdhdr."UDATE" AS "Invoice BP Last Change Date",
                cdhdr."UTIME" AS "Invoice BP Last Change Time",
                cdhdr."TCODE" AS "Invoice BP Last Change TCode",
                ROW_NUMBER() OVER (
                    PARTITION BY cdhdr."OBJECTID"
                    ORDER BY
                        cdhdr."UDATE" DESC,
                        cdhdr."UTIME" DESC,
                        cdhdr."CHANGENR" DESC
                ) AS rn
            FROM "SAPHANADB"."CDHDR" cdhdr
            WHERE cdhdr."MANDANT" = '{MANDT}'
              AND cdhdr."OBJECTCLAS" = 'DEBI'
              AND cdhdr."OBJECTID" IN ({bp_sql})
        )
        SELECT
            "Invoice-to BP",
            "Invoice BP Last Changed By",
            "Invoice BP Last Change Date",
            "Invoice BP Last Change Time",
            "Invoice BP Last Change TCode"
        FROM ranked_change
        WHERE rn = 1
        """
        frame = hana_query(sql)
        if not frame.empty:
            frames.append(frame)

    if not frames:
        return pd.DataFrame(columns=columns)

    result = pd.concat(frames, ignore_index=True).drop_duplicates(
        subset=["Invoice-to BP"], keep="first"
    )
    result = normalize_sap_text_columns(
        result,
        [
            "Invoice-to BP",
            "Invoice BP Last Changed By",
            "Invoice BP Last Change Time",
            "Invoice BP Last Change TCode",
        ],
    )
    result["Invoice BP Last Change Date"] = format_sap_date_series(
        result["Invoice BP Last Change Date"]
    )
    return ensure_columns(result, columns)[columns]


# ============================================================
# Current company warehouse stock
# ============================================================
def fetch_company_stock_locations(chassis_list: List[str]) -> pd.DataFrame:
    columns = [
        "Chassis",
        "Company Stock Current Location",
        "Stock Plant",
        "Stock Storage Location",
        "Stock Qty",
        "Stock Sales Order",
        "Stock Material",
    ]
    frames = []

    for batch in chunk_list(chassis_list):
        chassis_sql = sql_quote_list(batch)
        sql = f"""
        SELECT DISTINCT
            objk."SERNR" AS "Chassis",
            nsmka."WERKS" AS "Stock Plant",
            nsmka."LGORT" AS "Stock Storage Location",
            nsmka."KALAB" AS "Stock Qty",
            nsmka."VBELN" AS "Stock Sales Order",
            nsmka."MATNR" AS "Stock Material"
        FROM "SAPHANADB"."NSDM_V_MSKA" nsmka
        INNER JOIN "SAPHANADB"."SER02" ser02
            ON ser02."MANDT" = nsmka."MANDT"
           AND ser02."SDAUFNR" = nsmka."VBELN"
           AND ser02."POSNR" = '{SALES_ORDER_ITEM}'
        INNER JOIN "SAPHANADB"."OBJK" objk
            ON objk."MANDT" = ser02."MANDT"
           AND objk."OBKNR" = ser02."OBKNR"
        WHERE nsmka."MANDT" = '{MANDT}'
          AND objk."SERNR" IN ({chassis_sql})
          AND nsmka."KALAB" > 0
          AND (
                (nsmka."WERKS" = '3211'
                 AND nsmka."LGORT" IN (
                    '0001','0002','0003','0004',
                    '0005','0006','0007','0008'
                 ))
             OR (nsmka."WERKS" = '3411' AND nsmka."LGORT" = '0002')
          )
          AND (
                nsmka."MATNR" LIKE 'Z12%'
             OR nsmka."MATNR" LIKE 'Z19%'
          )
        """
        frame = hana_query(sql)
        if not frame.empty:
            frames.append(frame)

    if not frames:
        return pd.DataFrame(columns=columns)

    detail = pd.concat(frames, ignore_index=True).drop_duplicates()
    detail = normalize_sap_text_columns(
        detail,
        [
            "Chassis",
            "Stock Plant",
            "Stock Storage Location",
            "Stock Sales Order",
            "Stock Material",
        ],
    )
    detail["Warehouse Name"] = detail.apply(
        lambda row: WAREHOUSE_MAP.get(
            (row["Stock Plant"], row["Stock Storage Location"]),
            "Unknown Company Warehouse",
        ),
        axis=1,
    )

    detail["Plant/LGORT"] = (
        detail["Stock Plant"] + "/" + detail["Stock Storage Location"]
    )

    summary = (
        detail.sort_values(
            ["Chassis", "Warehouse Name", "Stock Plant", "Stock Storage Location"]
        )
        .groupby("Chassis", as_index=False)
        .agg(
            **{
                "Company Stock Current Location": (
                    "Warehouse Name",
                    lambda values: "; ".join(dict.fromkeys(values)),
                ),
                "Stock Plant": (
                    "Stock Plant",
                    lambda values: "; ".join(dict.fromkeys(values)),
                ),
                "Stock Storage Location": (
                    "Stock Storage Location",
                    lambda values: "; ".join(dict.fromkeys(values)),
                ),
                "Stock Qty": (
                    "Stock Qty",
                    lambda values: pd.to_numeric(values, errors="coerce").fillna(0).sum(),
                ),
                "Stock Sales Order": (
                    "Stock Sales Order",
                    lambda values: "; ".join(dict.fromkeys(values)),
                ),
                "Stock Material": (
                    "Stock Material",
                    lambda values: "; ".join(dict.fromkeys(values)),
                ),
            }
        )
    )
    return summary[columns]


# ============================================================
# SAP enrichment pipeline
# ============================================================
def build_sap_enrichment(firebase_df: pd.DataFrame):
    chassis_list = firebase_df["Chassis"].drop_duplicates().tolist()
    logger.info(
        "Firebase rows=%s, unique chassis=%s",
        len(firebase_df),
        len(chassis_list),
    )

    so_df = fetch_sales_orders_by_chassis(chassis_list)
    logger.info("Found chassis/SO rows: %s", len(so_df))

    sales_orders = (
        so_df["Sales Order SAP"]
        .replace("", pd.NA)
        .dropna()
        .drop_duplicates()
        .tolist()
    )

    pgi_df = fetch_last_valid_pgi(sales_orders)
    logger.info("Found PGI rows: %s", len(pgi_df))

    invoice_df = fetch_last_valid_invoice(sales_orders)
    logger.info("Found last valid invoice rows: %s", len(invoice_df))

    invoice_bps = (
        invoice_df["Invoice-to BP"]
        .replace("", pd.NA)
        .dropna()
        .drop_duplicates()
        .tolist()
        if not invoice_df.empty
        else []
    )

    try:
        bp_change_df = fetch_invoice_bp_last_change(invoice_bps)
        logger.info("Found Invoice-to BP change log rows: %s", len(bp_change_df))
    except Exception as error:
        logger.warning(
            "Invoice-to BP change log query failed; related columns will be blank: %s",
            error,
        )
        bp_change_df = pd.DataFrame(
            columns=[
                "Invoice-to BP",
                "Invoice BP Last Changed By",
                "Invoice BP Last Change Date",
                "Invoice BP Last Change Time",
                "Invoice BP Last Change TCode",
            ]
        )

    stock_summary_df = fetch_company_stock_locations(chassis_list)
    logger.info("Found company warehouse stock summary rows: %s", len(stock_summary_df))

    stock_detail_df = stock_summary_df.copy()

    result = firebase_df.merge(
        so_df,
        how="left",
        on="Chassis",
    )
    result = result.merge(
        pgi_df,
        how="left",
        on="Sales Order SAP",
    )
    result = result.merge(
        invoice_df,
        how="left",
        on="Sales Order SAP",
    )
    result = result.merge(
        bp_change_df,
        how="left",
        on="Invoice-to BP",
    )
    result = result.merge(
        stock_summary_df,
        how="left",
        on="Chassis",
    )

    text_columns = [
        "Sales Order SAP",
        "Sales Order Display",
        "Sales Order Item",
        "SO Material",
        "SO Material Name",
        "SO Is PGI",
        "SO PGI Post Date",
        "SO PGI Month",
        "SO PGI Material Document",
        "Last Invoice Number",
        "Last Invoice Date",
        "Invoice-to BP",
        "Invoice-to Name",
        "Invoice BP Last Changed By",
        "Invoice BP Last Change Date",
        "Invoice BP Last Change Time",
        "Invoice BP Last Change TCode",
        "Company Stock Current Location",
        "Stock Plant",
        "Stock Storage Location",
        "Stock Sales Order",
        "Stock Material",
    ]
    result = ensure_columns(result, text_columns)
    for column in text_columns:
        result[column] = result[column].fillna("").astype(str).str.strip()

    result["SO Is PGI"] = result["SO Is PGI"].replace("", "No_PGI")
    result.loc[result["SO Is PGI"] != "Yes", [
        "Last Invoice Number",
        "Last Invoice Date",
        "Invoice-to BP",
        "Invoice-to Name",
        "Invoice BP Last Changed By",
        "Invoice BP Last Change Date",
        "Invoice BP Last Change Time",
        "Invoice BP Last Change TCode",
    ]] = ""

    result["Company Stock Current Location"] = result[
        "Company Stock Current Location"
    ].replace("", "Not in company warehouse")

    if "Stock Qty" not in result.columns:
        result["Stock Qty"] = 0
    result["Stock Qty"] = pd.to_numeric(
        result["Stock Qty"], errors="coerce"
    ).fillna(0)

    final_columns = [
        "Firebase Key",
        "Chassis",
        "Duplicate Chassis in Firebase",
        "Transfer Current Location",
        "Transfer Target Location",
        "Transfer Type",
        "Transfer Saved At",
        "Sales Order SAP",
        "Sales Order Display",
        "Sales Order Item",
        "SO Material",
        "SO Material Name",
        "SO Is PGI",
        "SO PGI Post Date",
        "SO PGI Month",
        "SO PGI Material Document",
        "Last Invoice Number",
        "Last Invoice Date",
        "Invoice-to BP",
        "Invoice-to Name",
        "Invoice BP Last Changed By",
        "Invoice BP Last Change Date",
        "Invoice BP Last Change Time",
        "Invoice BP Last Change TCode",
        "Company Stock Current Location",
        "Stock Plant",
        "Stock Storage Location",
        "Stock Qty",
        "Stock Sales Order",
        "Stock Material",
    ]
    result = ensure_columns(result, final_columns)[final_columns]
    result = result.sort_values(
        ["Transfer Saved At", "Chassis", "Sales Order SAP"],
        na_position="last",
    )
    return result, stock_detail_df


# ============================================================
# Excel
# ============================================================
def write_excel(
    result_df: pd.DataFrame,
    firebase_df: pd.DataFrame,
    stock_detail_df: pd.DataFrame,
) -> None:
    report_summary = pd.DataFrame(
        {
            "Metric": [
                "Firebase Transfer Rows",
                "Unique Chassis",
                "Chassis with 3110 Sales Order",
                "Chassis PGI Yes",
                "Chassis No PGI",
                "Chassis in Company Warehouse",
                "Chassis Not in Company Warehouse",
                "Generated At",
            ],
            "Value": [
                len(firebase_df),
                firebase_df["Chassis"].nunique(),
                result_df.loc[
                    result_df["Sales Order SAP"].ne(""), "Chassis"
                ].nunique(),
                result_df.loc[
                    result_df["SO Is PGI"].eq("Yes"), "Chassis"
                ].nunique(),
                result_df.loc[
                    result_df["SO Is PGI"].eq("No_PGI"), "Chassis"
                ].nunique(),
                result_df.loc[
                    result_df["Company Stock Current Location"].ne(
                        "Not in company warehouse"
                    ),
                    "Chassis",
                ].nunique(),
                result_df.loc[
                    result_df["Company Stock Current Location"].eq(
                        "Not in company warehouse"
                    ),
                    "Chassis",
                ].nunique(),
                datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            ],
        }
    )

    with pd.ExcelWriter(OUTPUT_FILE, engine="openpyxl") as writer:
        result_df.to_excel(writer, sheet_name="Result", index=False)
        firebase_df.to_excel(writer, sheet_name="Firebase Raw", index=False)
        stock_detail_df.to_excel(writer, sheet_name="Warehouse Stock Detail", index=False)
        report_summary.to_excel(writer, sheet_name="Summary", index=False)

        header_fill = PatternFill(fill_type="solid", fgColor="1F4E78")
        header_font = Font(color="FFFFFF", bold=True)
        no_pgi_fill = PatternFill(fill_type="solid", fgColor="FFF2CC")
        outside_fill = PatternFill(fill_type="solid", fgColor="F4CCCC")

        for sheet_name, worksheet in writer.sheets.items():
            worksheet.freeze_panes = "A2"
            worksheet.auto_filter.ref = worksheet.dimensions

            for cell in worksheet[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center")

            for column_cells in worksheet.columns:
                column_letter = column_cells[0].column_letter
                max_length = 0
                for cell in column_cells:
                    value = "" if cell.value is None else str(cell.value)
                    max_length = max(max_length, min(len(value), 60))
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
                worksheet.column_dimensions[column_letter].width = max(
                    12,
                    min(max_length + 2, 45),
                )

        result_ws = writer.sheets["Result"]
        header_map = {cell.value: cell.column for cell in result_ws[1]}
        pgi_col = header_map.get("SO Is PGI")
        location_col = header_map.get("Company Stock Current Location")

        for row_number in range(2, result_ws.max_row + 1):
            if pgi_col and result_ws.cell(row_number, pgi_col).value == "No_PGI":
                result_ws.cell(row_number, pgi_col).fill = no_pgi_fill
            if (
                location_col
                and result_ws.cell(row_number, location_col).value
                == "Not in company warehouse"
            ):
                result_ws.cell(row_number, location_col).fill = outside_fill



# ============================================================
# Write SAP enrichment back to Firebase stock_transfer
# ============================================================
def update_stock_transfer_with_sap(result_df: pd.DataFrame) -> None:
    """
    Write selected SAP result fields back to each original Firebase record:

        /stock_transfer/{Firebase Key}/<field>

    The Firebase Key captured during the initial read is used, so records are
    updated in place even when the same chassis appears more than once.
    """
    if result_df.empty:
        logger.info("No result rows available for Firebase update.")
        return

    fields_to_write = [
        "Sales Order Display",
        "SO Is PGI",
        "SO PGI Post Date",
        "Last Invoice Number",
        "Last Invoice Date",
        "Invoice-to Name",
        "Invoice BP Last Changed By",
        "Invoice BP Last Change Date",
        "Company Stock Current Location",
    ]

    working = result_df.copy()
    for column in ["Firebase Key", "Sales Order SAP"] + fields_to_write:
        if column not in working.columns:
            working[column] = ""
        working[column] = working[column].fillna("").astype(str).str.strip()

    # A chassis can occasionally resolve to multiple 3110 sales orders. For a
    # single Firebase record, keep the numerically/latest-looking SAP SO row.
    working = working.sort_values(
        ["Firebase Key", "Sales Order SAP"],
        ascending=[True, False],
    ).drop_duplicates(subset=["Firebase Key"], keep="first")

    initialize_firebase()
    root_ref = db.reference("/")
    current_records = db.reference(f"/{FIREBASE_NODE}").get() or {}

    updates = {}
    audit_updates = {}
    synced_at = utc_now_iso()
    for _, row in working.iterrows():
        firebase_key = row["Firebase Key"]
        if not firebase_key:
            continue

        current_record = (
            current_records.get(firebase_key, {})
            if isinstance(current_records, dict)
            else {}
        )
        current_record = current_record if isinstance(current_record, dict) else {}
        changed_fields = {}
        snapshot_before = {}
        snapshot_after = {}

        for field in fields_to_write:
            value = row[field]
            previous_value = clean_text(current_record.get(field))
            snapshot_before[field] = previous_value
            snapshot_after[field] = value
            if previous_value != value:
                changed_fields[field] = {
                    "before": previous_value,
                    "after": value,
                }
            # Write an empty string when no value exists so stale prior values
            # are cleared rather than silently retained.
            updates[f"{FIREBASE_NODE}/{firebase_key}/{field}"] = value

        if changed_fields:
            audit_key = str(uuid.uuid4())
            audit_updates[f"{FIREBASE_AUDIT_NODE}/{firebase_key}/{audit_key}"] = {
                "action": "sap_enrichment_update",
                "transferId": firebase_key,
                "chassis": clean_text(row.get("Chassis")),
                "changedAt": synced_at,
                "changedBy": {
                    "type": "system",
                    "name": "Stock Transfer SAP Sync",
                    "email": "",
                    "company": "Local SAP sync worker",
                    "host": socket.gethostname(),
                },
                "source": "sap_sync",
                "changedFields": changed_fields,
                "snapshotBefore": snapshot_before,
                "snapshotAfter": snapshot_after,
            }

    if not updates:
        logger.info("No Firebase SAP enrichment fields to update.")
        return

    batch_size = 500
    update_items = list({**updates, **audit_updates}.items())
    for start in range(0, len(update_items), batch_size):
        batch = dict(update_items[start:start + batch_size])
        root_ref.update(batch)
        logger.info(
            "Firebase SAP enrichment/audit batch updated: %s paths",
            len(batch),
        )

    logger.info(
        "Firebase /%s updated for %s records with %s fields each; audit entries: %s.",
        FIREBASE_NODE,
        len(working),
        len(fields_to_write),
        len(audit_updates),
    )


def get_safe_firebase_key(value) -> str:
    return re.sub(r"[.#$\[\]/]", "_", str(value or ""))


def get_configured_base_url(config: dict) -> str:
    base_url = clean_text(config.get("appBaseUrl")) or DEFAULT_APP_BASE_URL
    return base_url.rstrip("/")


def has_sales_order(transfer: dict) -> bool:
    return bool(clean_text(transfer.get("Sales Order Display")))


def get_workflow(transfer: dict) -> dict:
    workflow = transfer.get("workflow")
    return workflow if isinstance(workflow, dict) else {}


def is_stock_transfer_email_active(transfer: dict) -> bool:
    workflow = get_workflow(transfer)
    return not (
        transfer.get("deletedAt")
        or transfer.get("cancelledAt")
        or workflow.get("purchaseDoneAt")
    )


def get_nsm_recipient(config: dict) -> str:
    recipients = config.get("recipients") if isinstance(config.get("recipients"), dict) else {}
    return clean_text(recipients.get("ceo"))


def get_nsm_cc_recipient(config: dict) -> str:
    cc_recipients = (
        config.get("ccRecipients")
        if isinstance(config.get("ccRecipients"), dict)
        else {}
    )
    return clean_text(cc_recipients.get("ceo"))


def get_stock_transfer_workflow_url(config: dict, transfer_id: str) -> str:
    query = urlencode({"taskTransfer": transfer_id})
    return f"{get_configured_base_url(config)}/#/stock-transfer-workflow/ceo?{query}"


def build_stock_transfer_task_email_html(transfer: dict, title: str, workflow_url: str) -> str:
    chassis = html.escape(clean_text(transfer.get("chassis")) or "-")
    current_location = html.escape(clean_text(transfer.get("currentLocation")) or "-")
    target_location = html.escape(clean_text(transfer.get("targetLocation")) or "-")
    escaped_url = html.escape(workflow_url)

    return f"""
    <div style="font-family:Arial,sans-serif;max-width:560px;margin:0 auto;color:#172554;">
      <div style="border-left:4px solid #2563eb;background:#eff6ff;padding:14px 16px;border-radius:10px;">
        <div style="font-weight:700;color:#1d4ed8;margin-bottom:4px;">Please confirm this stock transfer task.</div>
        <div style="color:#1e40af;font-size:13px;">Open the link and confirm once your part is done.</div>
      </div>
      <div style="margin-top:16px;border:1px solid #bfdbfe;border-radius:12px;overflow:hidden;">
        <div style="padding:10px 14px;border-bottom:1px solid #dbeafe;"><span style="display:inline-block;width:120px;color:#1d4ed8;font-weight:700;">Chassis</span><span style="color:#172554;">{chassis}</span></div>
        <div style="padding:10px 14px;border-bottom:1px solid #dbeafe;"><span style="display:inline-block;width:120px;color:#1d4ed8;font-weight:700;">From</span><span style="color:#172554;">{current_location}</span></div>
        <div style="padding:10px 14px;border-bottom:1px solid #dbeafe;"><span style="display:inline-block;width:120px;color:#1d4ed8;font-weight:700;">To</span><span style="color:#172554;">{target_location}</span></div>
        <div style="padding:10px 14px;"><span style="display:inline-block;width:120px;color:#1d4ed8;font-weight:700;">Current task</span><span style="color:#172554;">NSM Approval</span></div>
      </div>
      <div style="margin-top:16px;border:1px solid #bfdbfe;border-radius:12px;padding:14px 16px;background:#eff6ff;">
        <div style="font-weight:700;color:#1d4ed8;margin-bottom:8px;">Subtasks</div>
        <ul style="margin:0;padding-left:20px;color:#1e3a8a;line-height:1.55;">
          <li>Review the stock transfer request.</li>
          <li>Confirm NSM approval can proceed.</li>
        </ul>
      </div>
      <div style="margin-top:20px;text-align:center;">
        <a href="{escaped_url}" style="display:inline-block;background:#2563eb;color:#ffffff;text-decoration:none;padding:12px 22px;border-radius:999px;font-weight:700;">Open and Confirm</a>
      </div>
    </div>
    """


def should_queue_nsm_email(transfer: dict) -> bool:
    workflow = get_workflow(transfer)
    return (
        isinstance(transfer, dict)
        and is_stock_transfer_email_active(transfer)
        and has_sales_order(transfer)
        and not workflow.get("ceoApprovedAt")
        and not workflow.get("ceoEmailQueuedAt")
        and not workflow.get("ceoEmailJobId")
        and not workflow.get("ceoEmailSendingAt")
    )


def queue_nsm_email_jobs_after_sap_sync(result_df: pd.DataFrame) -> None:
    if result_df.empty or "Firebase Key" not in result_df.columns:
        return

    initialize_firebase()
    config = db.reference(f"/{FIREBASE_CONFIG_NODE}").get() or {}
    config = config if isinstance(config, dict) else {}
    recipient = get_nsm_recipient(config)
    cc_recipient = get_nsm_cc_recipient(config)
    unique_keys = [
        clean_text(firebase_key)
        for firebase_key in result_df["Firebase Key"].dropna().unique()
        if clean_text(firebase_key)
    ]

    queued_count = 0
    skipped_count = 0
    now_iso = utc_now_iso()

    for firebase_key in unique_keys:
        transfer_ref = db.reference(f"/{FIREBASE_NODE}/{firebase_key}")
        transfer = transfer_ref.get() or {}
        if not isinstance(transfer, dict) or not should_queue_nsm_email(transfer):
            skipped_count += 1
            continue

        if not recipient:
            transfer_ref.child("workflow").update(
                {
                    "ceoEmailError": "Missing recipient for ceo",
                    "ceoStatus": "Missing NSM approval recipient",
                }
            )
            skipped_count += 1
            continue

        workflow_url = get_stock_transfer_workflow_url(config, firebase_key)
        chassis = clean_text(transfer.get("chassis"))
        title = f"Action Required: Stock Transfer{f' {chassis}' if chassis else ''}"
        email_step = "nsm_approval"
        job_id = f"{get_safe_firebase_key(firebase_key)}_{email_step}"
        job_ref = db.reference(f"/{FIREBASE_EMAIL_JOBS_NODE}/{job_id}")
        existing_job = job_ref.get()

        if (
            isinstance(existing_job, dict)
            and clean_text(existing_job.get("status")).lower()
            in ACTIVE_EMAIL_JOB_STATUSES
        ):
            transfer_ref.child("workflow").update(
                {
                    "ceoEmailQueuedAt": now_iso,
                    "ceoEmailJobId": job_id,
                    "ceoEmailStep": email_step,
                    "ceoEmailRecipient": recipient,
                    "ceoEmailError": "",
                    "ceoStatus": "Pending NSM approval",
                }
            )
            skipped_count += 1
            continue

        job_ref.set(
            {
                "status": "pending",
                "step": email_step,
                "role": "ceo",
                "to": recipient,
                "cc": cc_recipient,
                "title": title,
                "content": build_stock_transfer_task_email_html(
                    transfer,
                    title,
                    workflow_url,
                ),
                "attempts": int(existing_job.get("attempts") or 0)
                if isinstance(existing_job, dict)
                else 0,
                "createdAt": now_iso,
                "updatedAt": now_iso,
                "lastError": None,
                "failedAt": None,
                "transferId": firebase_key,
                "source": "stock_transfer_sap_sync",
                "taskTransferId": firebase_key,
                "workflowUrl": workflow_url,
                "approveLink": workflow_url,
                "chassis": transfer.get("chassis") or "",
                "currentLocation": transfer.get("currentLocation") or "",
                "targetLocation": transfer.get("targetLocation") or "",
                "salesOrderDisplay": transfer.get("Sales Order Display") or "",
                "transferCategory": transfer.get("Stock Transfer Category") or "",
            }
        )
        transfer_ref.child("workflow").update(
            {
                "ceoEmailQueuedAt": now_iso,
                "ceoEmailJobId": job_id,
                "ceoEmailStep": email_step,
                "ceoEmailRecipient": recipient,
                "ceoEmailError": "",
                "ceoStatus": "Pending NSM approval",
            }
        )
        queued_count += 1

    logger.info(
        "SAP sync NSM email queue checked: queued=%s skipped=%s",
        queued_count,
        skipped_count,
    )


def get_multiple_sales_order_abnormal_rows(result_df: pd.DataFrame) -> pd.DataFrame:
    if result_df.empty:
        return pd.DataFrame()
    working = result_df.copy()
    for column in ["Firebase Key", "Sales Order SAP"]:
        if column not in working.columns:
            working[column] = ""
        working[column] = working[column].fillna("").astype(str).str.strip()

    so_counts = (
        working.loc[working["Sales Order SAP"].ne("")]
        .groupby("Firebase Key")["Sales Order SAP"]
        .nunique()
    )
    abnormal_keys = so_counts[so_counts > 1].index.tolist()
    if not abnormal_keys:
        return pd.DataFrame()
    return working[working["Firebase Key"].isin(abnormal_keys)].copy()


def write_abnormal_chassis(result_df: pd.DataFrame) -> None:
    abnormal_df = get_multiple_sales_order_abnormal_rows(result_df)
    initialize_firebase()
    abnormal_ref = db.reference(f"/{FIREBASE_ABNORMAL_NODE}")

    if abnormal_df.empty:
        for firebase_key in result_df.get("Firebase Key", pd.Series(dtype=str)).dropna().unique():
            if clean_text(firebase_key):
                abnormal_ref.child(clean_text(firebase_key)).delete()
        return

    detail_columns = [
        "Sales Order SAP",
        "Sales Order Display",
        "Sales Order Item",
        "SO Material",
        "SO Material Name",
        "SO Is PGI",
        "SO PGI Post Date",
        "Last Invoice Number",
        "Last Invoice Date",
        "Invoice-to BP",
        "Invoice-to Name",
        "Company Stock Current Location",
        "Stock Plant",
        "Stock Storage Location",
        "Stock Qty",
        "Stock Sales Order",
        "Stock Material",
    ]

    for firebase_key, group in abnormal_df.groupby("Firebase Key"):
        candidates = []
        for _, row in group.iterrows():
            candidate = {}
            for column in detail_columns:
                value = row[column] if column in row.index else ""
                candidate[column] = "" if pd.isna(value) else str(value)
            candidates.append(candidate)

        first_row = group.iloc[0]
        abnormal_ref.child(clean_text(firebase_key)).set(
            {
                "reason": "MULTIPLE_SALES_ORDERS",
                "firebaseKey": clean_text(firebase_key),
                "chassis": clean_text(first_row.get("Chassis")),
                "currentLocation": clean_text(
                    first_row.get("Transfer Current Location")
                ),
                "targetLocation": clean_text(
                    first_row.get("Transfer Target Location")
                ),
                "detectedAt": utc_now_iso(),
                "candidateSalesOrders": candidates,
                "recommendedAction": "Review manually before relying on the selected sales order.",
            }
        )

    logger.info("Abnormal chassis records written: %s", abnormal_df["Firebase Key"].nunique())


def mark_sync_done(firebase_keys: List[str]) -> None:
    if not firebase_keys:
        return
    node_ref = db.reference(f"/{FIREBASE_NODE}")
    updates = {}
    synced_at = utc_now_iso()
    for firebase_key in firebase_keys:
        updates[f"{firebase_key}/sapSyncStatus"] = "done"
        updates[f"{firebase_key}/sapSyncedAt"] = synced_at
        updates[f"{firebase_key}/sapSyncError"] = ""
        updates[f"{firebase_key}/sapSyncRetryAfter"] = ""
        updates[f"{firebase_key}/sapSyncClaimId"] = ""
    node_ref.update(updates)


def mark_sync_error(firebase_keys: List[str], error: Exception) -> None:
    if not firebase_keys:
        return
    node_ref = db.reference(f"/{FIREBASE_NODE}")
    retry_after = (utc_now() + timedelta(seconds=ERROR_RETRY_SECONDS)).replace(
        microsecond=0
    ).isoformat().replace("+00:00", "Z")
    message = str(error)[:1000]
    updates = {}
    for firebase_key in firebase_keys:
        updates[f"{firebase_key}/sapSyncStatus"] = "error"
        updates[f"{firebase_key}/sapSyncError"] = message
        updates[f"{firebase_key}/sapSyncRetryAfter"] = retry_after
        updates[f"{firebase_key}/sapSyncClaimId"] = ""
    node_ref.update(updates)


def process_pending_stock_transfer_once(limit: int = WORKER_BATCH_LIMIT) -> int:
    worker_id = get_worker_id()
    if not acquire_worker_lock(worker_id):
        logger.info("Another SAP sync worker is active; skipping this cycle.")
        return 0

    claimed_keys = []
    try:
        records_for_sync = fetch_stock_transfer_records_for_sync(limit)
        if not records_for_sync:
            logger.info("No stock transfer records available for SAP sync.")
            return 0

        for firebase_key in records_for_sync:
            if claim_stock_transfer_record(firebase_key, worker_id):
                claimed_keys.append(firebase_key)

        if not claimed_keys:
            logger.info("No stock transfer records claimed for SAP sync.")
            return 0

        claimed_records = {
            firebase_key: records_for_sync[firebase_key]
            for firebase_key in claimed_keys
            if firebase_key in records_for_sync
        }
        firebase_df = stock_transfer_records_to_df(claimed_records)
        if firebase_df.empty:
            mark_sync_done(claimed_keys)
            return 0

        result, _stock_detail_df = build_sap_enrichment(firebase_df)
        write_abnormal_chassis(result)
        update_stock_transfer_with_sap(result)
        queue_nsm_email_jobs_after_sap_sync(result)
        mark_sync_done(claimed_keys)
        logger.info("SAP sync worker processed records: %s", len(claimed_keys))
        return len(claimed_keys)
    except Exception as error:
        logger.exception("SAP sync worker cycle failed: %s", error)
        mark_sync_error(claimed_keys, error)
        return 0
    finally:
        release_worker_lock(worker_id)


def run_worker(interval_seconds: int, limit: int) -> None:
    logger.info(
        "Starting stock transfer SAP sync worker: interval=%ss limit=%s",
        interval_seconds,
        limit,
    )
    while True:
        process_pending_stock_transfer_once(limit=limit)
        time.sleep(interval_seconds)

# ============================================================
# Main
# ============================================================
def main() -> None:
    parser = argparse.ArgumentParser(
        description="Regent RV stock_transfer SAP enrichment"
    )
    parser.add_argument(
        "--worker",
        action="store_true",
        help="Run forever as a Firebase pending queue worker.",
    )
    parser.add_argument(
        "--once",
        action="store_true",
        help="Process one pending worker batch and exit.",
    )
    parser.add_argument(
        "--interval",
        type=int,
        default=WORKER_INTERVAL_SECONDS,
        help="Worker polling interval in seconds.",
    )
    parser.add_argument(
        "--limit",
        type=int,
        default=WORKER_BATCH_LIMIT,
        help="Maximum stock transfer records per worker cycle; use 0 to sync all eligible records.",
    )
    args = parser.parse_args()

    if args.worker:
        run_worker(interval_seconds=args.interval, limit=args.limit)
        return

    if args.once:
        run_log_file = add_once_run_log_handler()
        logger.info(
            "Starting stock_transfer SAP sync once: limit=%s log_file=%s",
            args.limit,
            run_log_file,
        )
        print(f"Run log: {run_log_file}")
        processed = process_pending_stock_transfer_once(limit=args.limit)
        logger.info(
            "Finished stock_transfer SAP sync once: processed=%s",
            processed,
        )
        print(f"Done. Stock transfer SAP sync records processed: {processed}")
        return

    logger.info("Starting Firebase stock_transfer -> SAP report")

    firebase_df = fetch_stock_transfer()
    if firebase_df.empty:
        raise RuntimeError("Firebase /stock_transfer has no chassis records.")

    result, stock_detail_df = build_sap_enrichment(firebase_df)
    write_abnormal_chassis(result)
    write_excel(result, firebase_df, stock_detail_df)
    update_stock_transfer_with_sap(result)
    queue_nsm_email_jobs_after_sap_sync(result)

    logger.info("Excel generated: %s", OUTPUT_FILE)
    print(f"Done. Excel generated: {OUTPUT_FILE}")
    print(f"Firebase /{FIREBASE_NODE} records updated with SAP fields.")


if __name__ == "__main__":
    try:
        main()
    except RefreshError as error:
        error_report = format_firebase_auth_error(error)
        logger.error("%s", error_report)
        raise SystemExit(1) from error
    except FirebaseCredentialError as error:
        logger.error("%s", error)
        raise SystemExit(1) from error
